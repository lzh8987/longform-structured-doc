# -*- coding: utf-8 -*-
"""
assemble_docx.py — 把「有序 Markdown 章节 + 附录」打包成合规 Word 文档。

基于 python-docx 生成，确保 Word/WPS 都能正常打开。

特性：
  - 自动映射 Markdown 标题：#→H1 … ######→H6（均为 Word 原生 Heading 样式，
    进导航窗格大纲、可被「引用→目录」自动抓取）。
  - 标题自动编号（默认开启）：通过**原生多级列表编号（numPr 指向 multilevel 列表）**
    实现——编号由 Word 按位置自动计算，移动/增删章节时编号与目录**自动跟随重排**，
    无需手工改号。可用 --no-numbering 关闭（标题仅保留层级样式、不带编号）。
    >>> 编号「样式」可配置：--numbering 1|2|3|4|5|6 选模板，或用 --custom-numbering 内联 JSON 自定义，
        每一级可独立指定 numFmt（decimal / chineseCounting …）与 lvlText（如 "第%1章"）。
  - 排版（字体/字号/行距/段距）可配置：--typo 1|2|3|4 选模板，或 --custom-typo 内联 JSON 自定义。
        支持正文/标题字体、标题字号（各级独立）、正文号、行距倍数、段前/段后（磅）、首行缩进（字符）。
  - 可选 --toc 生成「目录域」：
      · 在 settings.xml 写入 updateFields=true 并将域标 dirty，Word/WPS 打开即自动刷新，
        无需手动“更新域”，且能算出真实页码；
      · 同时在域的缓存结果里预填由 markdown 算出的编号目录作兜底（编号格式跟随所选模板），
        即使在不刷新的查看器（部分 WPS/预览窗格/邮件预览）中也能看到完整目录，
        不会裸奔出占位符。
  - 引用块(>) 作为缩进段落；代码块(```) 转为等宽段落；
    表格行(|...|) 转为真 Word 表格（带边框、表头加灰底）。
  - 自动抽取全文【待补充 / 待确认 / 需核实】生成「附录B 空缺事项汇总」（按三类分列）。
  - 可选 --page-numbers 生成页脚页码（域，随 updateFields 自动刷新）。
  - 内嵌 A4 页边距；中文字体默认由模板决定（正文/标题分开）。

用法：
  python assemble_docx.py --title "标题" --subtitle "副标题" --author "编制组" \
      --chapters ch1.md ch2.md --out out.docx --toc --page-numbers \
      --numbering 1 --typo 1

  # 自定编号（例：1 / 1.1 / 1.1.1 / （4））
  --custom-numbering '{"levels":[{"fmt":"decimal","text":"%1","suff":"space"},
      {"fmt":"decimal","text":"%1.%2","suff":"space"},
      {"fmt":"decimal","text":"%1.%2.%3","suff":"space"},
      {"fmt":"decimal","text":"（%4）","suff":"space"}]}'

  # 自定排版（例：正文小四宋体、1.5 倍行距、段后 6 磅、首行缩进 2 字）
  --custom-typo '{"body_font":"宋体","heading_font":"黑体","body_size":12,
      "line_spacing":1.5,"para_after":6,"first_line_indent":2}'

  # 或写 manifest.json（numbering / typography 字段可为模板 id 字符串或对象）：
  python assemble_docx.py --manifest manifest.json
"""
import os
import re
import json
import argparse
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_PATH = os.path.join(SCRIPT_DIR, "templates.json")

# 每轮运行时的配置（单进程单次运行，模块级即可）
CFG = {}

# 排版默认（custom-typo 缺字段时回退，也用于模板合并基）
DEFAULT_TYPO = {
    "body_font": "宋体",
    "heading_font": "黑体",
    "ascii_font": "Times New Roman",
    "title_size": 22,
    "subtitle_size": 14,
    "h_size": [16, 14, 12, 11, 10, 10],
    "body_size": 12,
    "line_spacing": 1.5,
    "para_before": 0,
    "para_after": 6,
    "first_line_indent": 2,
}


def load_templates():
    try:
        with open(TEMPLATES_PATH, encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {"numbering": {}, "typography": {}}


def set_run_font(run, font_name, size, bold=False, ascii_font=None):
    run.font.name = ascii_font or font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), font_name)
    rfonts.set(qn('w:ascii'), ascii_font or font_name)
    rfonts.set(qn('w:hAnsi'), ascii_font or font_name)
    run.font.size = Pt(size)
    run.font.bold = bold


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def _fmt_para(p, is_heading=False):
    """把当前排版配置应用到段落格式（行距/段距/首行缩进）。"""
    t = CFG['typo']
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = t['line_spacing']
    if is_heading:
        pf.space_before = Pt(t.get('heading_before', t['para_before']))
        pf.space_after = Pt(t.get('heading_after', t['para_after']))
        pf.first_line_indent = None
    else:
        pf.space_before = Pt(t['para_before'])
        pf.space_after = Pt(t['para_after'])
        if t.get('first_line_indent', 0):
            pf.first_line_indent = Pt(t['first_line_indent'] * t['body_size'])


def int_to_chinese(n):
    """阿拉伯数字 → 中文小写（一二三…），用于 numbering 模板的 chineseCounting 兜底渲染。"""
    if n <= 0:
        return str(n)
    if n < 10:
        return "零一二三四五六七八九"[n]
    if n < 20:
        return "十" + ("零一二三四五六七八九"[n % 10] if n % 10 else "")
    if n < 100:
        t, o = divmod(n, 10)
        return "零一二三四五六七八九"[t] + "十" + ("零一二三四五六七八九"[o] if o else "")
    # 100 ~ 999
    s = ""
    for i, ch in enumerate(reversed(str(n))):
        d = int(ch)
        units = ["", "十", "百", "千"]
        if d:
            s = "零一二三四五六七八九"[d] + units[i] + s
    return s


def circled(n):
    """圈码数字 ①..⑳，对应 numbering numFmt=decimalEnclosedCircle 的兜底渲染。"""
    if 1 <= n <= 20:
        return chr(0x2460 + n - 1)
    return str(n)


def to_letter(n, upper=False):
    """阿拉伯数字 → 字母（1->a, 27->aa），对应 lowerLetter/upperLetter 兜底渲染。"""
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(ord('a') + r) + s
    return s.upper() if upper else s


class OutlineCounter:
    """仅用于预填目录兜底的编号推算；render 时按模板格式化每一级（支持到 6 级）。"""

    def __init__(self):
        self.c = self.s = self.u = self.v = self.w = self.x = 0

    def bump(self, level):
        if level == 1:
            self.c += 1
            self.s = self.u = self.v = self.w = self.x = 0
        elif level == 2:
            self.s += 1
            self.u = self.v = self.w = self.x = 0
        elif level == 3:
            self.u += 1
            self.v = self.w = self.x = 0
        elif level == 4:
            self.v += 1
            self.w = self.x = 0
        elif level == 5:
            self.w += 1
            self.x = 0
        elif level == 6:
            self.x += 1

    def counters(self, level):
        vals = [None, self.c, self.s, self.u, self.v, self.w, self.x]
        return {k: vals[k] for k in range(1, level + 1)}


def _render_level(level, counters):
    """按当前 numbering 模板把某级计数渲染成显示串（与 Word 渲染对齐）。"""
    levels = CFG['num_levels']
    disp = {}
    for k in range(1, level + 1):
        fmt = levels[k - 1]['fmt']
        val = counters[k]
        if fmt == 'chineseCounting':
            disp[k] = int_to_chinese(val)
        elif fmt == 'decimalEnclosedCircle':
            disp[k] = circled(val)
        elif fmt in ('lowerLetter', 'upperLetter'):
            disp[k] = to_letter(val, upper=(fmt == 'upperLetter'))
        else:
            disp[k] = str(val)
    text = levels[level - 1]['text']
    out = text
    for k in range(1, level + 1):
        out = out.replace('%%%d' % k, disp[k])
    out = re.sub(r'%[1-9]', '', out)
    return out


# 前导编号（md 标题里常自带），开启自动编号时统一剥掉，交给 Word 按位置重排，
# 避免「1.1 1.1 系统总体架构」式重复。
_LEAD_NUM_RE = re.compile(
    r'^\s*(?:'
    r'(?:\d+\.)+\d*\s*'                                   # 1.1.1 / 1.1
    r'|[（(]\s*\d+\s*[)）]\s*'                            # （1）/ (1)
    r'|第\s*[一二三四五六七八九十百千零]+\s*[章节目节]\s*'     # 第一章
    r'|[（(]\s*[一二三四五六七八九十百千零]+\s*[)）]\s*'      # （一）
    r'|[一二三四五六七八九十百千零]+\s*[、.．]\s*'            # 一、/ 一.
    r'|\d+\s*[、.．]\s*'                                  # 1、/ 1.
    r'|\d{1,3}\s+'                                       # 裸数字+空格（避开4位年份）
    r'|[①-⑳]\s*'                                         # ① 圈码
    r'|[a-zA-Z]\s*[.．、]\s*'                             # a. / a、
    r')'
)


def _strip_leading_number(text):
    return _LEAD_NUM_RE.sub('', text).strip()


def _iter_heading_lines(chapter_paths, max_level=6):
    for p in chapter_paths:
        with open(p, encoding="utf-8") as f:
            txt = f.read()
        for line in txt.split("\n"):
            line = line.rstrip()
            m = re.match(r"^(#{1,%d})\s+(.*)$" % max_level, line)
            if not m:
                continue
            lvl = len(m.group(1))
            if lvl > max_level:
                continue
            yield lvl, m.group(2).strip()


def _static_toc_lines(chapter_paths, max_level=6, numbering=True):
    """由 markdown 预先算出的目录文本行，作为目录域的缓存兜底（编号格式跟随模板）。"""
    counter = OutlineCounter() if numbering else None
    lines = []
    for lvl, raw in _iter_heading_lines(chapter_paths, max_level):
        title = _strip_leading_number(raw) if numbering else raw
        if numbering and counter:
            counter.bump(lvl)
            num = _render_level(lvl, counter.counters(lvl))
            lines.append(("  " * (lvl - 1)) + num + " " + title)
        else:
            lines.append(("  " * (lvl - 1)) + title)
    return lines or ["（暂无章节标题）"]


def add_toc(doc, chapter_paths, max_level=6, numbering=True):
    t = CFG['typo']
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("目录")
    set_run_font(run, t['heading_font'], t['h_size'][0], bold=True, ascii_font=t['ascii_font'])
    _fmt_para(p, is_heading=True)

    paragraph = doc.add_paragraph()
    run = paragraph.add_run()

    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    fld_char1.set(qn('w:dirty'), 'true')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = ' TOC \\o "1-%d" \\h \\z \\u ' % max_level
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'separate')

    static_lines = _static_toc_lines(chapter_paths, max_level, numbering=numbering)

    fld_char3 = OxmlElement('w:fldChar')
    fld_char3.set(qn('w:fldCharType'), 'end')

    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    for i, line in enumerate(static_lines):
        if i > 0:
            br = OxmlElement('w:br')
            run._r.append(br)
        tt = OxmlElement('w:t')
        tt.set(qn('xml:space'), 'preserve')
        tt.text = line
        run._r.append(tt)
    run._r.append(fld_char3)


def add_footer_page_numbers(doc):
    t = CFG['typo']
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = ' PAGE '
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'separate')
    placeholder = OxmlElement('w:t')
    placeholder.text = '1'
    fld_char3 = OxmlElement('w:fldChar')
    fld_char3.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run._r.append(placeholder)
    run._r.append(fld_char3)
    set_run_font(run, t['ascii_font'], 10)


def parse_md_table(blk):
    lines = [l.strip() for l in blk.strip().split("\n") if l.strip()]
    if len(lines) < 2:
        return None
    if not (lines[0].startswith("|") and lines[1].startswith("|")):
        return None

    def split_row(l):
        l = l.strip()
        if l.startswith("|"):
            l = l[1:]
        if l.endswith("|"):
            l = l[:-1]
        return [c.strip() for c in l.split("|")]

    sep = split_row(lines[1])
    if not all(re.match(r"^:?-{2,}:?$", c) for c in sep if c != ""):
        return None
    rows = [split_row(lines[0])]
    for l in lines[2:]:
        if l.startswith("|"):
            rows.append(split_row(l))
    return rows


def _add_numpr(paragraph, ilvl, num_id):
    """给段落挂原生列表编号（ilvl + numId），编号由 Word 按位置自动计算/重排。"""
    pPr = paragraph._p.get_or_add_pPr()
    old = pPr.find(qn('w:numPr'))
    if old is not None:
        pPr.remove(old)
    numPr = OxmlElement('w:numPr')
    i = OxmlElement('w:ilvl')
    i.set(qn('w:val'), str(ilvl))
    numPr.append(i)
    n = OxmlElement('w:numId')
    n.set(qn('w:val'), str(num_id))
    numPr.append(n)
    pPr.append(numPr)


def add_markdown_block(doc, blk, numbering=False, num_id=None):
    t = CFG['typo']
    blk = blk.strip("\n")
    if not blk.strip():
        return
    lines = blk.split("\n")
    first = lines[0]

    def titled(level, raw):
        text = _strip_leading_number(raw) if numbering else raw
        p = doc.add_heading(text, level=level)
        for run in p.runs:
            set_run_font(run, t['heading_font'], t['h_size'][level - 1], bold=True, ascii_font=t['ascii_font'])
        _fmt_para(p, is_heading=True)
        if numbering and num_id is not None:
            _add_numpr(p, level - 1, num_id)
        return p

    if first.startswith("# "):
        titled(1, first[2:].strip())
    elif first.startswith("## "):
        titled(2, first[3:].strip())
    elif first.startswith("### "):
        titled(3, first[4:].strip())
    elif first.startswith("#### "):
        titled(4, first[5:].strip())
    elif first.startswith("##### "):
        titled(5, first[6:].strip())
    elif first.startswith("###### "):
        titled(6, first[7:].strip())
    elif first.startswith(">"):
        joined = " ".join(l.lstrip("> ").strip() for l in lines)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        run = p.add_run(joined.strip())
        set_run_font(run, t['body_font'], t['body_size'], ascii_font=t['ascii_font'])
        _fmt_para(p, is_heading=False)
    elif first.startswith("```"):
        if blk.count("```") >= 2:
            end = blk.rfind("```")
            inner = blk[3:end].strip()
            for line in inner.split("\n"):
                p = doc.add_paragraph()
                run = p.add_run(line)
                set_run_font(run, "Courier New", 10)
                _fmt_para(p, is_heading=False)
    elif first.startswith("|"):
        rows = parse_md_table(blk)
        if rows:
            ncol = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncol)
            table.style = 'Table Grid'
            table_size = max(9, t['body_size'] - 2)
            for i, row in enumerate(table.rows):
                for j, cell in enumerate(row.cells):
                    text = rows[i][j] if j < len(rows[i]) else ""
                    cell.text = text
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            set_run_font(run, t['body_font'], table_size, bold=(i == 0), ascii_font=t['ascii_font'])
                    if i == 0:
                        set_cell_shading(cell, "D9D9D9")
        else:
            p = doc.add_paragraph()
            run = p.add_run(blk.replace("\n", "  |  "))
            set_run_font(run, t['body_font'], t['body_size'], ascii_font=t['ascii_font'])
            _fmt_para(p, is_heading=False)
    else:
        joined = "".join(lines)
        p = doc.add_paragraph()
        run = p.add_run(joined)
        set_run_font(run, t['body_font'], t['body_size'], ascii_font=t['ascii_font'])
        _fmt_para(p, is_heading=False)


def _ensure_numbering(doc, levels, max_level=6):
    """在 numbering.xml 中创建一个 multilevel 列表（按 levels 定义 1~max_level 级），
    返回其 numId；标题段落通过 numPr 引用它即获得原生自动编号。"""
    num_root = doc.part.numbering_part.element

    ids = []
    for a in num_root.findall(qn('w:abstractNum')):
        v = a.get(qn('w:abstractNumId'))
        if v is not None:
            try:
                ids.append(int(v))
            except ValueError:
                pass
    aid = (max(ids) + 1) if ids else 0

    nids = []
    for n in num_root.findall(qn('w:num')):
        v = n.get(qn('w:numId'))
        if v is not None:
            try:
                nids.append(int(v))
            except ValueError:
                pass
    nid = (max(nids) + 1) if nids else 1

    abs_num = OxmlElement('w:abstractNum')
    abs_num.set(qn('w:abstractNumId'), str(aid))
    abs_num.set(qn('w:multiLevelType'), 'multilevel')
    for i in range(max_level):
        lvl_def = levels[i] if i < len(levels) else {"fmt": "decimal", "text": '.'.join('%%%d' % (k + 1) for k in range(i + 1)), "suff": "space"}
        lvl = OxmlElement('w:lvl')
        lvl.set(qn('w:ilvl'), str(i))
        lvl.set(qn('w:tpl'), '0')
        start = OxmlElement('w:start')
        start.set(qn('w:val'), '1')
        lvl.append(start)
        fmt = OxmlElement('w:numFmt')
        fmt.set(qn('w:val'), lvl_def.get('fmt', 'decimal'))
        lvl.append(fmt)
        txt = OxmlElement('w:lvlText')
        txt.set(qn('w:val'), lvl_def.get('text', '%s' % (i + 1)))
        lvl.append(txt)
        suff = OxmlElement('w:suff')
        suff.set(qn('w:val'), lvl_def.get('suff', 'space'))
        lvl.append(suff)
        abs_num.append(lvl)
    num_root.append(abs_num)

    num = OxmlElement('w:num')
    num.set(qn('w:numId'), str(nid))
    an = OxmlElement('w:abstractNumId')
    an.set(qn('w:val'), str(aid))
    num.append(an)
    num_root.append(num)
    return nid


def _enable_update_fields_on_open(doc):
    """让 Word/WPS 打开文档时自动刷新所有域（目录、页码）。"""
    settings = doc.settings.element
    upd = settings.find(qn('w:updateFields'))
    if upd is None:
        upd = OxmlElement('w:updateFields')
        settings.append(upd)
    upd.set(qn('w:val'), 'true')


def build_docx(title, subtitle, author, chapter_paths, appendices, out_path,
               toc=False, page_numbers=False, numbering=True, max_level=6,
               num_levels=None, typo=None):
    CFG['typo'] = typo or dict(DEFAULT_TYPO)
    CFG['num_levels'] = num_levels or [
        {"fmt": "decimal", "text": "%1", "suff": "space"},
        {"fmt": "decimal", "text": "%1.%2", "suff": "space"},
        {"fmt": "decimal", "text": "%1.%2.%3", "suff": "space"},
        {"fmt": "decimal", "text": "%1.%2.%3.%4", "suff": "space"},
        {"fmt": "decimal", "text": "%1.%2.%3.%4.%5", "suff": "space"},
        {"fmt": "decimal", "text": "%1.%2.%3.%4.%5.%6", "suff": "space"},
    ]
    t = CFG['typo']

    doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    # Normal / Body Text 样式：字体与字号（段落级行距/段距在 _fmt_para 统一施加）
    styles = doc.styles
    for style_name in ['Normal', 'Body Text']:
        try:
            style = styles[style_name]
            style.font.name = t['ascii_font']
            style._element.rPr.rFonts.set(qn('w:eastAsia'), t['body_font'])
            style.font.size = Pt(t['body_size'])
        except KeyError:
            pass

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_run_font(run, t['heading_font'], t['title_size'], bold=True, ascii_font=t['ascii_font'])
    _fmt_para(p, is_heading=True)

    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(subtitle)
        set_run_font(run, t['body_font'], t['subtitle_size'], ascii_font=t['ascii_font'])
        _fmt_para(p, is_heading=True)

    # 先创建多级列表编号定义（返回 numId），供标题段落引用
    num_id = _ensure_numbering(doc, CFG['num_levels'], max_level) if numbering else None

    if toc:
        add_toc(doc, chapter_paths, max_level=max_level, numbering=numbering)

    for p in chapter_paths:
        with open(p, encoding="utf-8") as f:
            txt = f.read()
        for blk in re.split(r"\n\s*\n", txt):
            add_markdown_block(doc, blk, numbering=numbering, num_id=num_id)

    for ap_title, ap_lines in appendices:
        # 附录标题保持 Heading 1 样式（进大纲/目录），但不挂编号
        p = doc.add_heading(ap_title, level=1)
        for run in p.runs:
            set_run_font(run, t['heading_font'], t['h_size'][0], bold=True, ascii_font=t['ascii_font'])
        _fmt_para(p, is_heading=True)
        for ln in ap_lines:
            p = doc.add_paragraph()
            run = p.add_run(ln)
            set_run_font(run, t['body_font'], t['body_size'], ascii_font=t['ascii_font'])
            _fmt_para(p, is_heading=False)

    if page_numbers:
        add_footer_page_numbers(doc)

    _enable_update_fields_on_open(doc)

    doc.core_properties.title = title
    doc.core_properties.author = author or ""

    doc.save(out_path)
    print("[OK] docx written -> %s" % out_path)


TODO_PATTERNS = [
    ("待补充", r"【待补充[^】]{0,200}】"),
    ("待确认", r"【待确认[^】]{0,200}】"),
    ("需核实", r"【需核实[^】]{0,200}】"),
]


def collect_todo(chapter_paths):
    found = {kind: [] for kind, _ in TODO_PATTERNS}
    for p in chapter_paths:
        with open(p, encoding="utf-8") as f:
            txt = f.read()
        for kind, pat in TODO_PATTERNS:
            for m in re.finditer(pat, txt):
                tt = m.group(0)
                if tt not in found[kind]:
                    found[kind].append(tt)
    lines = []
    for kind, _ in TODO_PATTERNS:
        if found[kind]:
            lines.append("【%s】" % kind)
            lines.extend("• " + tt for tt in found[kind])
    return [("附录B 空缺事项汇总", lines or ["（无）"])]


def _resolve_numbering(args, manifest):
    templates = load_templates()
    custom = getattr(args, 'custom_numbering', None)
    if manifest and isinstance(manifest.get('numbering'), dict):
        custom = json.dumps(manifest['numbering'])
    if custom:
        obj = json.loads(custom) if isinstance(custom, str) else custom
        if 'levels' not in obj:
            raise SystemExit("--custom-numbering / manifest.numbering 必须包含 levels 字段")
        return obj['levels']
    # 模板 id：manifest 优先（仅当其为字符串），否则命令行
    nid = args.numbering
    if manifest and isinstance(manifest.get('numbering'), str):
        nid = manifest['numbering']
    preset = templates.get('numbering', {}).get(nid)
    if not preset:
        print("[WARN] numbering 模板 '%s' 不存在，回退到 '1'" % nid)
        preset = templates.get('numbering', {}).get('1')
    if not preset:
        return None
    return preset['levels']


def _resolve_typo(args, manifest):
    templates = load_templates()
    custom = getattr(args, 'custom_typo', None)
    if manifest and isinstance(manifest.get('typography'), dict):
        custom = json.dumps(manifest['typography'])
    if custom:
        obj = json.loads(custom) if isinstance(custom, str) else custom
        typo = dict(DEFAULT_TYPO)
        typo.update(obj)
        return typo
    tid = args.typo
    if manifest and isinstance(manifest.get('typography'), str):
        tid = manifest['typography']
    preset = templates.get('typography', {}).get(tid)
    if not preset:
        print("[WARN] typography 模板 '%s' 不存在，回退到 '1'" % tid)
        preset = templates.get('typography', {}).get('1')
    typo = dict(DEFAULT_TYPO)
    if preset:
        typo.update(preset)
    return typo


def main():
    ap = argparse.ArgumentParser(description="Markdown 章节 → Word docx 组装器（编号/排版可配置）")
    ap.add_argument("--manifest", help="manifest.json 路径（优先）")
    ap.add_argument("--title", default="文档标题")
    ap.add_argument("--subtitle", default="")
    ap.add_argument("--author", default="编制组")
    ap.add_argument("--chapters", nargs="+", help="有序章节 md 文件列表")
    ap.add_argument("--out", default="output.docx")
    ap.add_argument("--toc", action="store_true", help="插入目录域（打开即自动刷新，无需手动更新域）")
    ap.add_argument("--page-numbers", action="store_true", help="页脚插入页码")
    ap.add_argument("--no-numbering", action="store_true", help="关闭标题/目录自动编号（默认开启）")
    ap.add_argument("--max-toc-level", type=int, default=6, help="目录捕获层级 1~6（默认 6）")
    ap.add_argument("--numbering", default="1", help="编号模板 id：1|2|3|4|5|6（见 templates.json），默认 1")
    ap.add_argument("--typo", default="1", help="排版模板 id：1|2|3|4（见 templates.json），默认 1")
    ap.add_argument("--custom-numbering", help="内联 JSON 自定义编号（需含 levels 数组）")
    ap.add_argument("--custom-typo", help="内联 JSON 自定义排版（缺字段回退默认）")
    args = ap.parse_args()

    manifest = None
    if args.manifest:
        with open(args.manifest, encoding="utf-8") as f:
            manifest = json.load(f)

    title = (manifest or {}).get("title", args.title)
    subtitle = (manifest or {}).get("subtitle", args.subtitle)
    author = (manifest or {}).get("author", args.author)
    chapters = (manifest or {}).get("chapters", args.chapters or [])
    source_docs = (manifest or {}).get("source_docs", [])
    refs = (manifest or {}).get("refs", [])
    extra = (manifest or {}).get("extra_appendices", [])
    out = (manifest or {}).get("out", args.out)
    toc = (manifest or {}).get("toc", args.toc)
    page_numbers = (manifest or {}).get("page_numbers", args.page_numbers)
    numbering = (not (manifest or {}).get("no_numbering", False)) and (not args.no_numbering)
    max_level = max(1, min(6, (manifest or {}).get("max_toc_level", args.max_toc_level)))

    if not chapters:
        ap.error("必须提供 --chapters 或 manifest.chapters")

    num_levels = _resolve_numbering(args, manifest) if numbering else None
    typo = _resolve_typo(args, manifest)

    appendices = []
    if source_docs:
        appendices.append(("附录A 依据文件与资料清单", source_docs))
    appendices.extend(collect_todo(chapters))
    if refs:
        appendices.append(("附录C 参考文献", refs))
    appendices.extend((e["title"], e["lines"]) for e in extra)

    build_docx(title, subtitle, author, chapters, appendices, out,
               toc=toc, page_numbers=page_numbers,
               numbering=numbering, max_level=max_level,
               num_levels=num_levels, typo=typo)


if __name__ == "__main__":
    main()
